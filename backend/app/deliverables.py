"""Deliverable files: approval notes, Word, Excel and PowerPoint, written on the server (nothing leaves the plant).

Every file is marked as drafted by Yukti, advisory, and needing an accountable human's approval. It is stored under the
reports folder, fingerprinted (SHA-256), recorded in `reports` and in the audit log by the caller.
"""
from __future__ import annotations

import hashlib
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from .config import REPORTS
from .db import ex, new_id, now_iso

ADVISORY = "Drafted by Yukti (on-premise). Advisory only: it takes effect only when the accountable approver signs it."
MIME = {".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".pdf": "application/pdf", ".txt": "text/plain", ".py": "text/x-python", ".csv": "text/csv"}


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_")[:80] or "yukti"


def _path(stem: str, ext: str) -> Path:
    return REPORTS / f"{_safe(stem)}_{datetime.now():%Y%m%d-%H%M%S}_{new_id()[:6]}{ext}"


def record(path: Path, kind: str, user_id: str, title: str, task_id: str | None = None) -> dict[str, Any]:
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    rid = new_id()
    ex("INSERT INTO reports(id, kind, file_name, file_path, sha256, created_by, created_at, task_id, title) VALUES(?,?,?,?,?,?,?,?,?)",
       (rid, kind, path.name, str(path), sha, user_id, now_iso(), task_id, title[:200]))
    return {"id": rid, "kind": kind, "file_name": path.name, "url": f"/api/reports/{rid}/file", "sha256": sha, "title": title}


# ------------------------------------------------------------------ Word
def docx_file(title: str, sections: list[dict[str, Any]], stem: str | None = None, footer: str = ADVISORY) -> Path:
    """sections: [{heading, paragraphs: [str], bullets: [str], table: [[...], ...]}] (each key optional)."""
    import docx
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_heading(title, level=1)
    meta = d.add_paragraph()
    run = meta.add_run(f"Generated {datetime.now():%Y-%m-%d %H:%M} · Yukti Sovereign Industrial AI Workbench")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    for s in sections:
        if s.get("heading"):
            d.add_heading(s["heading"], level=2)
        for p in s.get("paragraphs") or []:
            d.add_paragraph(str(p))
        for b in s.get("bullets") or []:
            d.add_paragraph(str(b), style="List Bullet")
        for i, b in enumerate(s.get("steps") or [], 1):
            d.add_paragraph(f"{i}. {b}")
        rows = s.get("table") or []
        if rows:
            t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for ri, r in enumerate(rows):
                for ci, v in enumerate(r):
                    t.cell(ri, ci).text = "" if v is None else str(v)
                    if ri == 0:
                        for par in t.cell(ri, ci).paragraphs:
                            for rn in par.runs:
                                rn.bold = True
    sec = d.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.text = footer
    for rn in fp.runs:
        rn.font.size = Pt(8)
    out = _path(stem or title, ".docx")
    d.save(out)
    return out


# ------------------------------------------------------------------ Excel
def xlsx_file(title: str, sheets: list[dict[str, Any]], stem: str | None = None) -> Path:
    """sheets: [{name, header: [...], rows: [[...]], formulas: {"E2": "=C2*D2"}, notes: [str]}]."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for sh in sheets or [{"name": "Sheet1", "header": [], "rows": []}]:
        name = re.sub(r"[\[\]:*?/\\]+", " ", str(sh.get("name") or "Sheet")).strip()[:31] or "Sheet"
        base, n = name, 2
        while name.lower() in used:  # Excel sheet names are unique, case-insensitively
            name = f"{base[:27]} ({n})"
            n += 1
        used.add(name.lower())
        ws = wb.create_sheet(name)
        r0 = 1
        if sh.get("header"):
            ws.append(list(sh["header"]))
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="0E5A7A")
            r0 = 2
        for row in sh.get("rows") or []:
            ws.append(list(row))
        for cell, f in (sh.get("formulas") or {}).items():
            ws[cell] = f
        for i, n in enumerate(sh.get("notes") or [], start=ws.max_row + 2):
            ws.cell(row=i, column=1, value=str(n))
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
        ws.freeze_panes = ws.cell(row=r0, column=1)
    about = wb.create_sheet("About")
    about.append([title])
    about.append([ADVISORY])
    about.append([f"Generated {datetime.now():%Y-%m-%d %H:%M}"])
    out = _path(stem or title, ".xlsx")
    wb.save(out)
    return out


# ------------------------------------------------------------------ PowerPoint
def pptx_file(title: str, slides: list[dict[str, Any]], stem: str | None = None) -> Path:
    """slides: [{title, bullets: [str], notes: str, table: [[...]]}]. A title slide is added first."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = title
    s0.placeholders[1].text = f"Yukti · {datetime.now():%d %b %Y}\n{ADVISORY}"
    for sl in slides:
        rows = sl.get("table") or []
        layout = prs.slide_layouts[5 if rows else 1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = str(sl.get("title") or "")
        if rows:
            ncols = max(len(r) for r in rows)
            shape = s.shapes.add_table(len(rows), ncols, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.4) * len(rows))
            for ri, r in enumerate(rows):
                for ci in range(ncols):
                    cell = shape.table.cell(ri, ci)
                    cell.text = "" if ci >= len(r) or r[ci] is None else str(r[ci])
                    for p in cell.text_frame.paragraphs:
                        p.font.size = Pt(12)
        else:
            body = s.placeholders[1].text_frame
            body.clear()
            for i, b in enumerate(sl.get("bullets") or []):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = str(b)
                p.font.size = Pt(18)
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = str(sl["notes"])
    out = _path(stem or title, ".pptx")
    prs.save(out)
    return out


# ------------------------------------------------------------------ approval note for a finding
def approval_note(f: dict[str, Any], asset: dict[str, Any] | None, source: dict[str, Any] | None, procedure: dict[str, Any] | None,
                  approver: dict[str, Any] | None, lane: str, region: dict[str, Any] | None, verification: list[dict[str, Any]],
                  prepared_by: str) -> Path:
    a = asset or {}
    rows = [["Field", "Value"],
            ["Finding", f["title"]], ["Asset", f"{f['tag']} — {a.get('name') or ''}".strip(" —")],
            ["Unit / location", f"{a.get('unit') or ''} {a.get('location') or ''}".strip()],
            ["Discipline / approver lane", f"{(f.get('discipline') or '').title()} / {lane}"],
            ["Owner department", f.get("owner_department") or ""], ["Severity", f.get("severity") or ""],
            ["Due date", f.get("due_date") or ""], ["Approver", (approver or {}).get("display_name") or "to be assigned"]]
    ev = [f"Reading: {f.get('evidence') or ''}"]
    if source:
        ev.append(f"Source: {source.get('doc_number') or source.get('title')} rev {source.get('revision') or '-'}, page {f.get('page') or 1}")
    if region:
        ev.append(f"Evidence region on the page: x {region.get('x')}, y {region.get('y')}, width {region.get('w')}, height {region.get('h')} "
                  f"(fraction of the page), OCR confidence {region.get('confidence')}")
    checks = [["Check", "Result", "Detail"]] + [[v.get("check"), v.get("result"), v.get("detail")] for v in verification]
    sections = [
        {"heading": "1. Finding", "table": rows},
        {"heading": "2. Evidence", "bullets": ev},
        {"heading": "3. Verification against the asset master", "table": checks if len(checks) > 1 else [],
         "paragraphs": [] if len(checks) > 1 else ["No automated checks were recorded."]},
        {"heading": "4. Applicable procedure",
         "paragraphs": [f"{procedure.get('doc_number')} rev {procedure.get('revision')} — {procedure.get('title')}" if procedure
                        else "No procedure is linked to this asset in the knowledge base; the approver names one."]},
        {"heading": "5. Proposed action",
         "bullets": ["Confirm the reading in the field before acting.",
                     "Decide: repair / re-inspect / derate / accept with justification, within the due date.",
                     "Raise the work order and permit through the normal process (Yukti does not issue permits)."]},
        {"heading": "6. Sign-off", "table": [["Role", "Name", "Signature", "Date"], ["Prepared (Yukti draft for)", prepared_by, "", ""],
                                              ["Reviewed by (discipline)", "", "", ""], ["Approved by", (approver or {}).get("display_name") or "", "", ""]]},
    ]
    return docx_file(f"Approval note — {f['tag']}: {f['title'].replace('[EXAMPLE] ', '')}", sections, stem=f"approval_note_{f['tag']}")
