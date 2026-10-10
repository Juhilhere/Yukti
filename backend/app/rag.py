"""Ingestion (type sniff → per-page digital/scanned detection → OCR → chunk → index) and ABAC-filtered retrieval."""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import time
from pathlib import Path
from typing import Any

from . import audit
from .config import BLOBS
from .db import ex, j, new_id, now_iso, q, q1, uj
from .policy import Subject, pdp

_OCR = None
# pdfium is not thread-safe and uploads are ingested in worker threads: serialise PDF parsing/rendering and OCR.
import threading as _threading  # noqa: E402

_PDFIUM_LOCK = _threading.Lock()
_OCR_LOCK = _threading.Lock()


def ocr_engine():
    global _OCR
    if _OCR is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR = RapidOCR()
    return _OCR


def ocr_image(img) -> tuple[str, float]:
    import numpy as np
    arr = np.array(img.convert("RGB"))
    with _OCR_LOCK:
        res, _ = ocr_engine()(arr)
    if not res:
        return "", 0.0
    # sort boxes top-to-bottom, left-to-right, group into lines
    items = sorted(res, key=lambda r: (round(r[0][0][1] / 18), r[0][0][0]))
    lines: list[str] = []
    cur_y = None
    buf: list[str] = []
    for box, text, score in items:
        y = round(box[0][1] / 18)
        if cur_y is not None and y != cur_y:
            lines.append("  ".join(buf))
            buf = []
        buf.append(text)
        cur_y = y
    if buf:
        lines.append("  ".join(buf))
    conf = sum(float(r[2]) for r in res) / len(res)
    return "\n".join(lines), conf


# ------------------------------------------------------------------ jobs
def job_create(kind: str, document_id: str, user: str) -> str:
    jid = new_id()
    stages = [{"name": n, "status": "pending", "detail": ""} for n in
              ["Store & fingerprint", "Detect type & page modes", "Extract text / OCR", "Tag assets", "Chunk", "Index (BM25)"]]
    ex("INSERT INTO jobs(id, kind, status, stages_json, document_id, created_by, created_at, updated_at) VALUES(?,?,?,?,?,?,?,?)",
       (jid, kind, "queued", j(stages), document_id, user, now_iso(), now_iso()))
    return jid


def _stage(jid: str | None, idx: int, status: str, detail: str = "") -> None:
    if not jid:
        return
    r = q1("SELECT stages_json FROM jobs WHERE id=?", (jid,))
    st = uj(r["stages_json"], [])
    st[idx]["status"] = status
    st[idx]["detail"] = detail
    ex("UPDATE jobs SET stages_json=?, status=?, updated_at=? WHERE id=?", (j(st), "running", now_iso(), jid))


def _alias_map() -> dict[str, str]:
    return {r["alias"].upper(): r["tag"] for r in q("SELECT alias, tag FROM tag_alias")}


def detect_tags(text: str, amap: dict[str, str] | None = None) -> list[str]:
    amap = amap or _alias_map()
    found: set[str] = set()
    up = text.upper()
    for alias, tag in amap.items():
        if len(alias) <= 3:
            if re.search(rf"(?<![A-Z0-9-]){re.escape(alias)}(?![A-Z0-9])", up):
                found.add(tag)
        elif alias in up:
            found.add(tag)
    return sorted(found)


def _chunks(text: str, size: int = 900, overlap: int = 150) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) <= size:
        return [text] if text else []
    paras = re.split(r"\n\s*\n|\n(?=\d+[.)] )", text)
    out: list[str] = []
    buf = ""
    for p in paras:
        if len(buf) + len(p) + 1 > size and buf:
            out.append(buf.strip())
            buf = buf[-overlap:] + "\n" + p
        else:
            buf += "\n" + p
    if buf.strip():
        out.append(buf.strip())
    # hard split very long pieces
    final: list[str] = []
    for c in out:
        while len(c) > size * 1.6:
            final.append(c[: size])
            c = c[size - overlap:]
        final.append(c)
    return final


def extract(path: Path, mime_hint: str, jid: str | None) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".pdf":
        with _PDFIUM_LOCK:
            return _extract(path, mime_hint, jid)
    return _extract(path, mime_hint, jid)


def _extract(path: Path, mime_hint: str, jid: str | None) -> list[dict[str, Any]]:
    """Returns pages: [{page_no, mode, text, ocr_conf}]."""
    ext = path.suffix.lower()
    pages: list[dict[str, Any]] = []
    if ext == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(path))
        modes = []
        for i in range(len(pdf)):
            page = pdf[i]
            txt = page.get_textpage().get_text_range() or ""
            n_img = sum(1 for o in page.get_objects() if o.type == 3)  # FPDF_PAGEOBJ_IMAGE
            digital = len(txt.strip()) > 50
            modes.append("digital" if digital else "scanned")
            pages.append({"page_no": i + 1, "mode": "digital" if digital else "scanned", "text": txt, "ocr_conf": None,
                          "_img": n_img})
        _stage(jid, 1, "done", f"PDF · {len(pdf)} page(s): {modes.count('digital')} digital, {modes.count('scanned')} scanned")
        _stage(jid, 2, "running", "OCR on scanned pages" if "scanned" in modes else "Reading text layer")
        for p in pages:
            if p["mode"] == "scanned":
                img = pdf[p["page_no"] - 1].render(scale=2.2).to_pil()
                text, conf = ocr_image(img)
                p["text"], p["ocr_conf"] = text, round(conf, 3)
        n_ocr = sum(1 for p in pages if p["mode"] == "scanned")
        _stage(jid, 2, "done", f"OCR (RapidOCR PP-OCR ONNX) on {n_ocr} page(s)" if n_ocr else "Text layer extracted — no OCR needed")
    elif ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
        from PIL import Image
        _stage(jid, 1, "done", "Image · treated as a scanned page")
        _stage(jid, 2, "running", "OCR")
        im = Image.open(path)
        n_frames = getattr(im, "n_frames", 1)
        confs = []
        for fi in range(n_frames):
            im.seek(fi)
            text, conf = ocr_image(im.convert("RGB"))
            confs.append(conf)
            pages.append({"page_no": fi + 1, "mode": "image", "text": text, "ocr_conf": round(conf, 3)})
        _stage(jid, 2, "done", f"OCR on {n_frames} page(s), mean confidence {sum(confs) / len(confs):.2f}")
    elif ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        pages.append({"page_no": 1, "mode": "digital", "text": "\n".join(parts), "ocr_conf": None})
        _stage(jid, 1, "done", "DOCX · structured text")
        _stage(jid, 2, "done", f"{len(parts)} paragraphs/rows")
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        for si, ws in enumerate(wb.worksheets):
            rows = list(ws.iter_rows(values_only=True))
            # the header is the first row with at least two filled cells (exports often start with a title row)
            h0 = next((i for i, r in enumerate(rows) if sum(1 for v in r if v not in (None, "")) >= 2), None)
            if h0 is None:
                continue
            hdr = [str(h) if h is not None else "" for h in rows[h0]]
            lines = []
            for ri, r in enumerate(rows[h0 + 1:], start=h0 + 2):
                cells = [f"{hdr[i] if i < len(hdr) else i}={v}" for i, v in enumerate(r) if v not in (None, "")]
                if cells:
                    lines.append(f"[{ws.title}!{ri}] " + "; ".join(cells))
            pages.append({"page_no": si + 1, "mode": "digital", "text": "\n".join(lines), "ocr_conf": None})
        _stage(jid, 1, "done", f"XLSX · {len(pages)} sheet(s) as records")
        _stage(jid, 2, "done", "Rows converted to citable records")
    elif ext in TEXT_EXTS:
        raw = path.read_bytes()
        try:
            txt = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            txt = raw.decode("cp1252", errors="replace")
        pages.append({"page_no": 1, "mode": "digital", "text": txt, "ocr_conf": None})
        _stage(jid, 1, "done", "Plain text")
        _stage(jid, 2, "done", "Read")
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: " + ", ".join(sorted(SUPPORTED_EXTS)))
    for p in pages:
        p.pop("_img", None)
    return pages


TEXT_EXTS = {".txt", ".csv", ".md", ".log", ".json", ".xml"}
SUPPORTED_EXTS = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".docx", ".xlsx", ".xlsm"} | TEXT_EXTS


def ingest(document_id: str, jid: str | None, actor: str) -> None:
    doc = q1("SELECT * FROM documents WHERE id=?", (document_id,))
    path = Path(doc["file_path"])
    try:
        _stage(jid, 0, "done", f"sha256 {doc['sha256'][:16]}… · {doc['size_bytes'] // 1024} KB")
        _stage(jid, 1, "running")
        pages = extract(path, doc["mime"] or "", jid)
        amap = _alias_map()
        _stage(jid, 3, "running")
        all_text = "\n".join(p["text"] for p in pages)
        tags = sorted(set(uj(doc["asset_tags_json"], [])) | set(detect_tags(all_text, amap)))
        _stage(jid, 3, "done", ", ".join(tags[:12]) or "no asset tags found")
        _stage(jid, 4, "running")
        ex("DELETE FROM pages WHERE document_id=?", (document_id,))
        old = [r["id"] for r in q("SELECT id FROM chunks WHERE document_id=?", (document_id,))]
        for cid in old:
            ex("DELETE FROM chunks_fts WHERE chunk_id=?", (cid,))
        ex("DELETE FROM chunks WHERE document_id=?", (document_id,))
        n = 0
        for p in pages:
            ex("INSERT INTO pages(id, document_id, page_no, mode, text, ocr_conf) VALUES(?,?,?,?,?,?)",
               (new_id(), document_id, p["page_no"], p["mode"], p["text"], p["ocr_conf"]))
            for i, ch in enumerate(_chunks(p["text"])):
                cid = new_id()
                ctags = " ".join(detect_tags(ch, amap))
                ex("INSERT INTO chunks(id, document_id, page, ordinal, text, tags) VALUES(?,?,?,?,?,?)",
                   (cid, document_id, p["page_no"], i, ch, ctags))
                header = f"{doc['title']} {doc['doc_number'] or ''} rev {doc['revision'] or ''}"
                ex("INSERT INTO chunks_fts(text, tags, chunk_id) VALUES(?,?,?)", (header + "\n" + ch, ctags + " " + " ".join(tags), cid))
                n += 1
        if n == 0:
            raise ValueError("No readable text was found in this file (empty, image-only with unreadable scan, or a "
                             "spreadsheet without data rows). It is stored but cannot be searched or cited.")
        _stage(jid, 4, "done", f"{n} chunks")
        _stage(jid, 5, "done", "FTS5 BM25 index updated")
        ex("UPDATE documents SET pages=?, asset_tags_json=? WHERE id=?", (len(pages), j(tags), document_id))
        from . import graph, vectors
        graph.index_document(document_id)  # links to assets, procedures and referenced documents
        if vectors.ready():
            import threading
            threading.Thread(target=vectors.index_missing, daemon=True, name="embed-index").start()
        if jid:
            ex("UPDATE jobs SET status='done', updated_at=? WHERE id=?", (now_iso(), jid))
        audit.write(actor, "document.ingested", f"document:{document_id}",
                    {"pages": len(pages), "chunks": n, "scanned": sum(1 for p in pages if p["mode"] != "digital")})
    except Exception as e:  # pragma: no cover
        if jid:
            ex("UPDATE jobs SET status='error', error=?, updated_at=? WHERE id=?", (str(e), now_iso(), jid))
        audit.write(actor, "document.ingest_failed", f"document:{document_id}", {"error": str(e)})
        raise


def store_file(src: Path | bytes, file_name: str) -> tuple[Path, str, int]:
    data = src.read_bytes() if isinstance(src, Path) else src
    sha = hashlib.sha256(data).hexdigest()
    dest = BLOBS / sha[:2] / f"{sha}{Path(file_name).suffix.lower()}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    if not dest.exists():
        dest.write_bytes(data)
    return dest, sha, len(data)


def create_document(meta: dict[str, Any], file_name: str, data: bytes | Path, actor: str) -> str:
    path, sha, size = store_file(data, file_name)
    did = new_id()
    ext = Path(file_name).suffix.lower()
    mime = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".txt": "text/plain",
            ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}.get(ext, "application/octet-stream")
    # supersede older revision of the same document number
    if meta.get("doc_number") and meta.get("status", "CURRENT") == "CURRENT":
        for old in q("SELECT id, revision FROM documents WHERE doc_number=? AND status='CURRENT' AND department=?",
                     (meta["doc_number"], meta.get("department"))):
            if str(old["revision"]) != str(meta.get("revision")) and _rev_key(old["revision"]) < _rev_key(meta.get("revision")):
                ex("UPDATE documents SET status='SUPERSEDED' WHERE id=?", (old["id"],))
    ex("""INSERT INTO documents(id, title, doc_number, revision, status, doc_type, department, classification, effective_date,
          supersedes, asset_tags_json, file_path, file_name, mime, sha256, size_bytes, uploaded_by, created_at, is_example, is_public, source_url)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
       (did, meta.get("title") or Path(file_name).stem, meta.get("doc_number"), str(meta.get("revision") or ""),
        meta.get("status", "CURRENT"), meta.get("doc_type", "other"), meta.get("department", "Operations"),
        int(meta.get("classification", 1)), meta.get("effective_date"), meta.get("supersedes"),
        j(meta.get("asset_tags", [])), str(path), file_name, mime, sha, size, actor, now_iso(),
        int(bool(meta.get("is_example"))), int(bool(meta.get("is_public"))), meta.get("source_url")))
    return did


def _rev_key(rev: Any) -> tuple:
    """Natural ordering of revision labels: R9 < R10, rev B < rev C, 2 < 10."""
    parts = re.findall(r"\d+|[A-Za-z]+", str(rev or ""))
    return tuple((0, int(x)) if x.isdigit() else (1, x.upper()) for x in parts)


# ------------------------------------------------------------------ access helpers
def _units_for(tags: list[str]) -> list[str]:
    if not tags:
        return []
    marks = ",".join("?" * len(tags))
    return sorted({r["unit"] for r in q(f"SELECT unit FROM assets WHERE tag IN ({marks})", tags) if r["unit"]})


def doc_resource(d: dict[str, Any]) -> dict[str, Any]:
    tags = uj(d.get("asset_tags_json"), [])
    return {"type": "document", "id": d["id"], "department": d["department"], "classification": int(d["classification"]),
            "doc_type": d["doc_type"], "asset_tags": tags, "units": _units_for(tags), "status": d["status"]}


def can_read(subject: Subject, d: dict[str, Any], action: str = "read") -> bool:
    return pdp.decide(subject, action, doc_resource(d)).allowed


def visible_documents(subject: Subject) -> list[dict[str, Any]]:
    return [d for d in q("SELECT * FROM documents ORDER BY created_at DESC") if can_read(subject, d)]


# ------------------------------------------------------------------ retrieval
STOP = set("a an the of to in on for and or is are was were be what which who how when where why me my i we our you your "
           "give show tell please with from at by it this that these those do does did can could should would about any all "
           "need needs get list find".split())


# common Hindi / Kannada filler words: they would match almost every Hindi or Kannada document
STOP_INDIC = {"है", "हैं", "का", "की", "के", "को", "में", "से", "पर", "और", "या", "यह", "वह", "क्या", "कैसे", "कौन", "कब", "एक", "भी", "तो",
              "ही", "नहीं", "था", "थी", "हो", "कर", "करें", "करना", "दो", "दें", "बताओ", "बताएं", "मुझे", "हम", "आप",
              "ಮತ್ತು", "ಅಥವಾ", "ಈ", "ಆ", "ಏನು", "ಹೇಗೆ", "ಯಾವ", "ಯಾರು", "ಯಾವಾಗ", "ಒಂದು", "ಇದೆ", "ಇವೆ", "ಅಲ್ಲ", "ನನಗೆ", "ನೀವು", "ಮಾಡಿ",
              "ಮಾಡುವುದು", "ತಿಳಿಸಿ", "ಬಗ್ಗೆ", "ಅನ್ನು", "ಅನ್ನ", "ಗೆ", "ಲ್ಲಿ"}


def _fts_query(text: str, extra_tags: list[str]) -> str:
    # Latin words and tags, plus Hindi (Devanagari) and Kannada words, so documents written in those languages are found too
    toks = [t for t in re.findall(r"[A-Za-z0-9][A-Za-z0-9\-_.]*|[ऀ-ॿ]+|[ಀ-೿]+", text)
            if t.lower() not in STOP and t not in STOP_INDIC and len(t) > 1]
    toks = toks[:24] + extra_tags
    return " OR ".join('"' + t.replace('"', "") + '"' for t in dict.fromkeys(toks)) or '""'


# document types that answer a given kind of question best (metadata channel of the hybrid search)
INTENT_DOC_TYPES = {
    "procedure_lookup": {"SOP", "troubleshooting_guide"},
    "asset_dossier": {"datasheet", "drawing_sld", "drawing_pid", "work_order_export", "SOP"},
    "finding_review": {"inspection_report", "calibration_certificate"},
    "process_chemistry": {"MSDS", "process_manual"},
    "asset_expiry": {"calibration_certificate", "inspection_report", "asset_register"},
}
_RECORD = re.compile(r"\[([^\[\]!]{1,40}![0-9]{1,7})\]")


def retrieve(subject: Subject, text: str, k: int = 8, intent: str | None = None, embedder: Any = None) -> dict[str, Any]:
    """Hybrid, policy-filtered retrieval.

    Channels: BM25 keywords (FTS5), asset tags, document metadata (type that fits the question), knowledge-graph links
    (documents linked to the asset, procedures, referenced documents) and, when an embedding model runs, dense vectors.
    Every candidate document is checked with the policy engine (`cite`) before anything about it is used; withheld
    documents are only counted. Each source carries its revision, page or record ID, timestamp, confidence and the
    channels that found it."""
    from . import graph, vectors
    tags = detect_tags(text)
    fq = _fts_query(text, tags)
    t0 = time.time()
    sql = """SELECT c.id cid, c.document_id, c.page, c.text, c.tags, bm25(chunks_fts) score
             FROM chunks_fts JOIN chunks c ON c.id = chunks_fts.chunk_id
             WHERE chunks_fts MATCH ? ORDER BY score LIMIT 80"""
    try:
        rows = q(sql, (fq,))
    except Exception as e:  # noqa: BLE001 - never silently answer without evidence: retry with plain quoted terms
        from .engine import log
        # no question words in the log (the database's own message can quote them too): only how many terms there were
        log(f"[search] query with {fq.count(' OR ') + 1} term(s) failed ({type(e).__name__}); retrying with plain terms")
        words = [w for w in re.findall(r"[A-Za-z0-9][A-Za-z0-9\-]{1,}", text)][:24]
        try:
            rows = q(sql, (" OR ".join(f'"{w}"' for w in words),)) if words else []
        except Exception as e2:  # noqa: BLE001
            log(f"[search] fallback query with {len(words)} term(s) failed too ({type(e2).__name__})")
            rows = []
    channels: dict[str, set[str]] = {r["cid"]: {"bm25"} for r in rows}
    if tags:
        for r in rows:
            if any(t in (r["tags"] or "").split() for t in tags):
                channels[r["cid"]].add("tags")
    # dense vectors (only when an embedding model runs)
    dense = dict(vectors.search(text, 40, embedder)) if (embedder is not None or vectors.ready()) else {}
    extra_ids = [cid for cid in dense if cid not in channels]
    if extra_ids:
        marks = ",".join("?" * len(extra_ids))
        for r in q(f"SELECT c.id cid, c.document_id, c.page, c.text, c.tags, 0.0 score FROM chunks c WHERE c.id IN ({marks})", extra_ids):
            rows.append(r)
            channels[r["cid"]] = set()
    for cid in dense:
        if cid in channels:
            channels[cid].add("vector")
    # knowledge graph: documents linked to the asset (its procedures, drawings, referenced documents)
    linked = graph.linked_documents(tags) if tags else {}
    have_docs = {r["document_id"] for r in rows}
    for did in [d for d in linked if d not in have_docs][:6]:
        best = q("SELECT id cid, document_id, page, text, tags, 0.0 score FROM chunks WHERE document_id=? ORDER BY "
                 "(CASE WHEN " + " OR ".join(["tags LIKE ?"] * len(tags)) + " THEN 0 ELSE 1 END), page, ordinal LIMIT 1",
                 (did, *[f"%{t}%" for t in tags]))
        for r in best:
            rows.append(r)
            channels[r["cid"]] = set()
    for r in rows:
        if r["document_id"] in linked:
            channels[r["cid"]].add("graph")
    want_types = INTENT_DOC_TYPES.get(intent or "", set())

    docs: dict[str, dict[str, Any]] = {}
    allowed: dict[str, bool] = {}
    denied_docs: dict[str, dict[str, Any]] = {}
    base = [-float(r["score"]) for r in rows if "bm25" in channels[r["cid"]]]
    top = max(base) if base else 1.0
    scored: list[tuple[float, dict[str, Any]]] = []
    for r in rows:
        did = r["document_id"]
        if did not in docs:
            docs[did] = q1("SELECT * FROM documents WHERE id=?", (did,))
            allowed[did] = can_read(subject, docs[did], "cite")
        d = docs[did]
        if not allowed[did]:
            # company-level briefings are simply not for this person: no "withheld - request access" notice for them
            if d["doc_type"] != "company_briefing":
                denied_docs[did] = d
            continue
        ch = channels[r["cid"]]
        s = -float(r["score"]) if "bm25" in ch else 0.0
        if "vector" in ch:
            s += 0.5 * top * max(0.0, dense[r["cid"]])
        if not ch & {"bm25", "vector"}:  # reached only through the graph: below solid keyword hits, but in the evidence
            s = 0.35 * top
        if "tags" in ch:
            s *= 1.5
        if "graph" in ch and ch & {"bm25", "vector"}:
            s *= 1.15
        if want_types and d["doc_type"] in want_types:
            ch.add("metadata")
            s *= 1.15
        if d["status"] == "SUPERSEDED":
            s *= 0.8
        scored.append((s, {**r, "doc": d}))
    scored.sort(key=lambda x: -x[0])
    best_s = scored[0][0] if scored else 1.0
    out: list[dict[str, Any]] = []
    per_doc: dict[str, int] = {}
    ocr: dict[tuple[str, int], float | None] = {}
    for s, r in scored:
        did = r["document_id"]
        if per_doc.get(did, 0) >= 2:
            continue
        per_doc[did] = per_doc.get(did, 0) + 1
        d = r["doc"]
        key = (did, int(r["page"] or 0))
        if key not in ocr:
            pg = q1("SELECT mode, ocr_conf FROM pages WHERE document_id=? AND page_no=?", key)
            ocr[key] = float(pg["ocr_conf"]) if pg and pg["mode"] != "digital" and pg["ocr_conf"] is not None else None
        conf = max(0.0, min(1.0, s / best_s if best_s > 0 else 0.0)) * (ocr[key] if ocr[key] is not None else 1.0)
        rec = _RECORD.search(r["text"] or "")
        out.append({
            "id": f"S{len(out) + 1}", "document_id": did, "title": d["title"], "doc_number": d["doc_number"] or "",
            "revision": d["revision"] or "", "status": d["status"], "page": r["page"], "snippet": r["text"][:600],
            "record_id": rec.group(1) if rec else None,
            "locator": (f"record {rec.group(1)}" if rec else f"p.{r['page']}"),
            "score": round(s, 3), "confidence": round(conf, 2), "ocr_confidence": ocr[key],
            "channels": sorted(channels[r["cid"]]), "graph_reason": (linked.get(did) or [None])[0],
            "classification": ["PUBLIC", "INTERNAL", "RESTRICTED", "CONFIDENTIAL", "SECRET"][int(d["classification"])],
            "department": d["department"], "doc_type": d["doc_type"], "effective_date": d["effective_date"],
            "timestamp": d["effective_date"] or (d["created_at"] or "")[:10], "ingested_at": d["created_at"],
            "_text": r["text"],
        })
        if len(out) >= k:
            break
    denied_depts = sorted({d["department"] for d in denied_docs.values()}) if subject.clearance >= 1 else []
    return {"sources": out, "denied": {"count": len(denied_docs), "departments": denied_depts, "doc_types": []},
            "tags": tags, "related_tags": graph.related_tags(tags) if tags else [],
            "channels": {"bm25": True, "tags": bool(tags), "graph": bool(linked), "vector": bool(dense), "metadata": bool(want_types)},
            "latency_ms": round((time.time() - t0) * 1000, 1)}


# ------------------------------------------------------------------ dossier facts (KNOWN / MISSING / CONFLICTING)
def dossier_facts(subject: Subject, tag: str) -> list[dict[str, Any]]:
    from .scope import record_readable
    asset = q1("SELECT * FROM assets WHERE tag=?", (tag,))
    # structured plant records (asset master, CMMS, contacts) follow the same department / unit / clearance policy as documents
    records_ok = record_readable(subject, asset)
    wo_ok = record_readable(subject, asset, "work_order_export")
    rows = q("SELECT * FROM facts WHERE asset_tag=? ORDER BY slot, attribute", (tag,))
    groups: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for r in rows:
        groups.setdefault((r["slot"], r["attribute"]), []).append(r)
    facts: list[dict[str, Any]] = []
    for (slot, attr), cands in groups.items():
        visible = []
        withheld = 0
        for c in cands:
            d = q1("SELECT * FROM documents WHERE doc_number=? AND revision=?", (c["source_doc_number"], c["revision"]))
            if (d and not can_read(subject, d, "cite")) or (not d and not records_ok):  # fail closed
                withheld += 1
                continue
            visible.append((c, d))
        vals = [(c, d) for c, d in visible if c["value"] not in (None, "", "null")]
        if not vals:
            named = sorted({c["source_doc_number"] or c["source_kind"] for c, _ in visible if c["source_doc_number"] or c["source_kind"]})
            note = f"No authorised source holds '{attr}'."
            if named:
                note += " Expected in: " + ", ".join(named) + "."
            if withheld:
                note += f" {withheld} source(s) outside your access may hold it; you can request access."
            elif not named:
                note += " Expected in: wiring diagram / asset master."
            facts.append({"slot": slot, "attribute": attr, "value": None, "status": "MISSING", "candidates": [], "note": note})
            continue
        norm = {re.sub(r"\s+", "", str(c["value"]).lower()) for c, _ in vals}

        def rank(cd: tuple[dict[str, Any], Any]) -> tuple[int, str]:
            c, d = cd
            cur = 1 if (d is None or d["status"] == "CURRENT") else 0
            pri = {"datasheet": 3, "sop": 3, "drawing": 2, "asset_master": 1, "cmms": 2}.get(c["source_kind"] or "", 1)
            return (cur * 10 + pri, c["effective_date"] or "")

        best = max(vals, key=rank)
        cands_out = [{"value": f"{c['value']}{(' ' + c['unit']) if c['unit'] else ''}",
                      "source_label": f"{c['source_doc_number'] or c['source_kind']}" + (f" rev {c['revision']}" if c["revision"] else ""),
                      "source_id": (d["id"] if d else ""), "revision": c["revision"] or "",
                      "effective": c["effective_date"] or "", "status": d["status"] if d else "RECORD",
                      "recommended": (c is best[0]) and len(norm) > 1} for c, d in vals]
        status = "CONFLICTING" if len(norm) > 1 else "KNOWN"
        bc = best[0]
        note = ""
        if status == "CONFLICTING":
            note = (f"Sources disagree. Recommended: {bc['source_doc_number'] or bc['source_kind']} rev {bc['revision']} "
                    "(latest CURRENT revision of the highest-priority source). Verify in field before acting.")
        facts.append({"slot": slot, "attribute": attr, "value": f"{bc['value']}{(' ' + bc['unit']) if bc['unit'] else ''}",
                      "status": status, "candidates": cands_out, "note": note})
    # structured CMMS facts
    wos = q("SELECT * FROM work_orders WHERE tag=? ORDER BY opened_at DESC", (tag,)) if wo_ok else []
    if wos:
        open_ = [w for w in wos if w["status"] == "OPEN"]
        bd = [w for w in wos if w["type"] == "BD"]
        facts.append({"slot": "history", "attribute": "open_work_orders", "value": ", ".join(w["wo_no"] for w in open_) or "none",
                      "status": "KNOWN", "candidates": [{"value": str(len(open_)), "source_label": "CMMS work orders", "source_id": "",
                                                          "revision": "", "effective": "", "status": "RECORD", "recommended": False}], "note": ""})
        causes: dict[str, int] = {}
        for w in bd:
            causes[w["cause"] or w["failure_code"] or "?"] = causes.get(w["cause"] or w["failure_code"] or "?", 0) + 1
        facts.append({"slot": "history", "attribute": "breakdowns_5y", "value": f"{len(bd)} ({', '.join(f'{k}×{v}' for k, v in causes.items())})",
                      "status": "KNOWN", "candidates": [], "note": ""})
    oc = q("SELECT * FROM contacts WHERE on_call=1 AND dept IN ('Electrical', 'Mechanical', 'Operations') ORDER BY dept")         if records_ok and subject.clearance >= 1 else []
    if oc:
        facts.append({"slot": "contacts", "attribute": "on_call", "value": "; ".join(f"{c['name']} ({c['dept']}, ext {c['ext']})" for c in oc[:3]),
                      "status": "KNOWN", "candidates": [], "note": ""})
    order = {"ratings": 0, "wiring": 1, "isolation": 2, "procedure": 3, "history": 4, "contacts": 5}
    facts.sort(key=lambda f: (order.get(f["slot"], 9), f["attribute"]))
    return facts
