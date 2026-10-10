"""Workbench features: durable work queue and workers, hot KV slots and chat checkpoints,
hybrid retrieval (graph, vectors, metadata) with full citations, the asset ledger, approver lanes, scan to finding with
human review, the agentic workbench and its sandbox, the isolated research layer, the forecast + optimizer plan and
Performance Engine profiles. No GPU or model needed."""
from __future__ import annotations

from datetime import date, timedelta

import numpy as np
import pytest


# ------------------------------------------------------------------ Laya -> queue -> worker
def test_laya_decision_selects_a_worker(login):
    a = login("admin")
    d = a.post("/api/laya/classify", json={"text": "which certificates expire in the next 30 days"}).json()
    assert d["worker"] == "analytics"
    assert a.post("/api/laya/classify", json={"text": "write a python script to plot A2 vibration trend"}).json()["worker"] == "coder"


def test_select_worker_rules():
    from app.workqueue import select_worker
    assert select_worker({"intent": "asset_dossier"}, True, "") == "ocr_vlm"
    assert select_worker({"intent": "asset_dossier"}, False, "why does A2 keep tripping? root cause please") == "reasoning"
    assert select_worker({"intent": "procedure_lookup"}, False, "isolate A2") == "small_llm"
    assert select_worker({"intent": "production_scenario"}, False, "") == "analytics"


def test_chat_turn_is_a_task_with_checkpoint_rules(login):
    u = login("ravi.e")
    ev = u.chat("A2 tripped - ratings, wiring, isolation and history")
    t = ev["task"]
    assert t["worker"] in ("small_llm", "reasoning") and t["label"]
    task = u.get(f"/api/tasks/{t['id']}").json()
    assert task["kind"] == "chat" and task["status"] == "error"  # no model in CI: recorded, not lost
    assert login("meera.me").get(f"/api/tasks/{t['id']}").status_code == 404  # tasks are private to their owner
    seen = login("admin").get(f"/api/tasks/{t['id']}").json()  # the queue overseer sees the shape of the work only
    assert seen["result"] == {} and seen["title"] == "chat" and all(set(x) <= {"name", "status"} for x in seen["steps"])


def test_policy_answers_are_tasks_on_the_policy_worker(login):
    u = login("ravi.e")
    ev = u.chat("write a poem about love")
    t = u.get(f"/api/tasks/{ev['task']['id']}").json()
    assert t["worker"] == "policy" and t["status"] == "done" and t["result"]["stop_reason"] == "off_topic"


def test_queue_recovers_background_work_after_restart(app_client):
    from app import workqueue
    from app.db import ex, q1
    tid = workqueue.create("workbench", "analytics", None, payload={}, status="running")
    cid = workqueue.create("chat", "small_llm", None, payload={}, status="running")
    ex("UPDATE tasks SET attempts=1 WHERE id=?", (tid,))
    workqueue.recover()
    assert q1("SELECT status FROM tasks WHERE id=?", (cid,))["status"] == "interrupted"
    st = q1("SELECT status FROM tasks WHERE id=?", (tid,))["status"]
    assert st in ("queued", "running", "done", "error")  # requeued (and possibly picked up again)
    ex("UPDATE tasks SET status='error' WHERE id=?", (tid,))


def test_admin_queue_view_hides_questions(login):
    a = login("admin")
    s = a.get("/api/admin/queue").json()
    assert {w["id"] for w in s["workers"]} >= {"policy", "small_llm", "ocr_vlm", "analytics", "reasoning", "coder", "workflow"}
    assert all(t["title"] == t["kind"] for t in s["recent"])
    assert login("ravi.e").get("/api/admin/queue").status_code == 403


# ------------------------------------------------------------------ hot KV (TTL + LRU) and checkpoints
def test_hot_kv_slots_are_sticky_lru_and_expire():
    from app.kvcache import HotKV
    kv = HotKV()
    kv.reset(2)
    a = kv.acquire("chat-a", 600)
    kv.release(a["slot"])
    b = kv.acquire("chat-b", 600)
    kv.release(b["slot"])
    assert a["slot"] != b["slot"] and a["state"] == "cold"
    again = kv.acquire("chat-a", 600)
    assert again == {"slot": a["slot"], "state": "hot", "evicted": None}
    kv.release(again["slot"])
    c = kv.acquire("chat-c", 600)  # both slots held: the least recently used chat (b) gives its slot up
    assert c["evicted"] == "chat-b" and c["slot"] == b["slot"]
    kv.release(c["slot"])
    kv.table["chat-a"]["at"] -= 10_000  # idle longer than the TTL: no longer hot
    assert kv.acquire("chat-a", 600)["state"] == "cold"


def test_busy_slots_are_never_shared():
    from app.kvcache import HotKV
    kv = HotKV()
    kv.reset(1)
    a = kv.acquire("x", 600)
    assert kv.acquire("y", 600)["slot"] is None  # the only slot is answering: no stickiness, llama.cpp picks
    kv.release(a["slot"])


def test_checkpoints_are_per_chat_and_private(login, app_client):
    from app import kvcache
    u = login("ravi.e")
    chat = u.post("/api/chats", json={}).json()
    kvcache.checkpoint(chat["id"], "m1", None, "small_llm", "gemma", {"slot": 0, "state": "hot", "cached_tokens": 812},
                       [{"tool": "search_documents"}], [{"rule": "G-GENERAL", "decision": "allow"}], ["S1"])
    cps = u.get(f"/api/chats/{chat['id']}/checkpoints").json()
    assert cps[0]["kv"]["cached_tokens"] == 812 and cps[0]["citations"] == ["S1"]
    assert login("meera.me").get(f"/api/chats/{chat['id']}/checkpoints").status_code == 404


def test_spilled_slot_files_follow_ttl_and_size_cap(tmp_path, monkeypatch):
    import os
    import time
    from app import kvcache
    monkeypatch.setattr(kvcache, "SLOT_DIR", tmp_path)
    for i, age in enumerate((10, 20, 5000)):
        p = tmp_path / f"chat-{i}.bin"
        p.write_bytes(b"x" * 2**20)
        os.utime(p, (time.time() - age, time.time() - age))
    assert kvcache.sweep_spilled(ttl=1000, max_mb=1.5) == 2  # one too old, then the least recently used over the cap
    assert [p.name for p in tmp_path.iterdir()] == ["chat-0.bin"]


# ------------------------------------------------------------------ hybrid retrieval and citations
def test_sources_carry_revision_locator_timestamp_confidence_and_channels(login):
    src = login("ravi.e").chat("A2 tripped - wiring and isolation per SOP-EL-014")["retrieval"]["sources"]
    assert src
    for s in src:
        assert {"revision", "locator", "timestamp", "confidence", "channels"} <= set(s)
        assert 0 <= s["confidence"] <= 1 and "bm25" in s["channels"] or "graph" in s["channels"] or "vector" in s["channels"]
    assert any("graph" in s["channels"] for s in src)
    assert any(s["record_id"] for s in login("ravi.e").chat("A2 open work orders history")["retrieval"]["sources"]
               if s["doc_number"] == "MNT-WO-A2")


def test_graph_links_assets_procedures_and_references(app_client):
    from app import graph
    linked = graph.linked_documents(["A2"])
    from app.db import q1
    sop = q1("SELECT id FROM documents WHERE doc_number='SOP-EL-014' AND status='CURRENT'")
    assert sop["id"] in linked
    assert any(e["rel"] == "procedure" for e in graph.neighbours(graph.asset_node("A2")))
    assert graph.stats()["by_relation"].get("supersedes", 0) >= 1


def test_graph_never_bypasses_policy(login):
    """A finance document linked to A2 in the graph still never reaches an electrician."""
    ev = login("ravi.e").chat("A2 boiler fuel cost audit findings")
    assert all(s["department"] != "Finance & Accounts" for s in ev["retrieval"]["sources"])


def test_dense_vectors_are_fused_and_still_policy_filtered(login, app_client):
    from app import rag, vectors
    from app.auth import subject_for
    from app.db import q1

    def fake(texts):  # deterministic "embedding": bag of a few plant words
        vocab = ["isolat", "seal", "audit", "thickness", "amine", "trip"]
        return [[t.lower().count(w) + 0.01 for w in vocab] for t in texts]

    assert vectors.index_missing(fake) > 0
    ravi = subject_for(q1("SELECT * FROM users WHERE username='ravi.e'"))
    ret = rag.retrieve(ravi, "seal leak on the booster pump", embedder=fake)
    assert ret["channels"]["vector"] and any("vector" in s["channels"] for s in ret["sources"])
    assert all(rag.can_read(ravi, q1("SELECT * FROM documents WHERE id=?", (s["document_id"],)), "cite") for s in ret["sources"])


def test_intent_metadata_boosts_matching_document_types(app_client):
    from app import rag
    from app.auth import subject_for
    from app.db import q1
    s = subject_for(q1("SELECT * FROM users WHERE username='ravi.e'"))
    out = rag.retrieve(s, "procedure to isolate A2 motor", intent="procedure_lookup")
    assert out["sources"][0]["doc_type"] in ("SOP", "troubleshooting_guide") and "metadata" in out["sources"][0]["channels"]


# ------------------------------------------------------------------ asset ledger
def test_ledger_has_extinguishers_licences_and_category_lead_times(login):
    p = login("priya.hse", fresh=True).get("/api/assets").json()
    cats = {i["category"] for a in p for i in a["items"]}
    assert {"fire_extinguisher", "licence", "calibration", "pressure_vessel_inspection"} <= cats
    pv = next(i for a in p for i in a["items"] if i["category"] == "pressure_vessel_inspection")
    assert pv["alert_days"] == 60


def test_register_asset_and_ledger_entry_with_interval(login):
    h = login("sunita.hse")  # HOD HSE: assets.edit for her department
    r = h.post("/api/assets", json={"tag": "FE-901", "name": "CO2 extinguisher, control room", "unit": "CDU-1", "class": "extinguisher",
                                    "vendor": "Example Safety", "serial": "EX-1"})
    assert r.status_code == 200 and r.json()["owner_department"] == "Health, Safety & Environment"
    issued = (date.today() - timedelta(days=350)).isoformat()
    it = h.post("/api/assets/FE-901/ledger", json={"category": "fire_extinguisher", "issued_on": issued, "ref_no": "HT-77",
                                                   "certificate": "IS 15683", "warranty_until": (date.today() + timedelta(days=400)).isoformat()})
    assert it.status_code == 200, it.text
    it = it.json()
    assert it["interval_months"] == 12 and it["status"] == "amber" and it["days_left"] <= 30
    assert any(a["tag"] == "FE-901" for a in h.get("/api/alerts").json())
    ren = h.post(f"/api/ledger/{it['id']}/renew", json={"issued_on": date.today().isoformat()}).json()
    assert ren["status"] == "green"
    assert login("ravi.e").post("/api/assets", json={"tag": "X-1", "name": "x"}).status_code == 403
    assert h.post("/api/assets/FE-901/ledger", json={"category": "fire_extinguisher"}).status_code == 422


def test_analytics_worker_answers_expiry_from_the_ledger(login):
    ev = login("priya.hse", fresh=True).chat("which certificates expire in the next 30 days")
    assert ev["task"]["worker"] == "analytics"
    assert ev["facts"]["tag"] == "ledger" or any(f["slot"] == "ledger" for f in ev["facts"]["facts"])


# ------------------------------------------------------------------ approver lanes, escalation, follow-up
def test_lane_approvers_and_escalation_chain(app_client):
    from app import lanes
    from app.db import q1
    uid = lambda n: q1("SELECT id FROM users WHERE username=?", (n,))["id"]  # noqa: E731
    assert lanes.approver_for("mechanical") == uid("rajesh.mm")
    assert lanes.approver_for("finance") == uid("kavita.fm")
    assert lanes.approver_for("hse") == uid("sunita.hse")  # no approver:hse holder: the department HOD
    assert lanes.next_approver(uid("rajesh.mm"), "Mechanical Maintenance") == uid("ramesh.pm")
    assert lanes.next_approver(uid("ramesh.pm"), "Mechanical Maintenance") == uid("director.md")


def test_follow_up_escalates_unacknowledged_and_reminds_overdue(app_client):
    from app import lanes
    from app.db import ex, j, new_id, q1
    fid = new_id()
    rajesh = q1("SELECT id FROM users WHERE username='rajesh.mm'")["id"]
    ex("""INSERT INTO findings(id, title, tag, discipline, severity, state, due_date, evidence, approver_id, history_json, created_at,
          owner_department, lane) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
       (fid, "test overdue", "V-101", "Mechanical", "high", "PENDING", (date.today() - timedelta(days=2)).isoformat(), "x", rajesh,
        j([]), "2020-01-01T00:00:00+00:00", "Mechanical Maintenance", "mechanical"))
    out = lanes.follow_up()
    f = q1("SELECT * FROM findings WHERE id=?", (fid,))
    assert out["escalated"] >= 1 and f["state"] == "ESCALATED" and f["approver_id"] != rajesh
    assert lanes.follow_up()["reminded"] == 0 or q1("SELECT 1 FROM findings WHERE id=?", (fid,))  # escalation was today's follow-up
    ex("UPDATE findings SET last_follow_up_at='2020-01-01T00:00:00+00:00' WHERE id=?", (fid,))
    assert lanes.follow_up()["reminded"] >= 1  # overdue: a reminder to the approver and the HOD ...
    before = q1("SELECT last_follow_up_at FROM findings WHERE id=?", (fid,))["last_follow_up_at"]
    lanes.follow_up()
    assert q1("SELECT last_follow_up_at FROM findings WHERE id=?", (fid,))["last_follow_up_at"] == before  # ... at most daily
    ex("UPDATE findings SET state='REJECTED' WHERE id=?", (fid,))


# ------------------------------------------------------------------ scan to finding
def _doc(login, number):
    return next(d for d in login("ramesh.pm").get("/api/documents").json() if d["doc_number"] == number)


def test_scan_fields_and_evidence_regions_from_ocr(app_client):
    from app import scan_findings
    from app.db import q1
    d = q1("SELECT * FROM documents WHERE doc_number='UT-E310-2026-07'")
    (n, img), = scan_findings.page_images(d)
    rows = scan_findings.ocr_rows(img)
    ex = scan_findings.extract_fields(rows, n)
    assert ex["fields"]["min_required_mm"]["value"] == "6.5" and ex["fields"]["inspection_date"]["value"] == "2026-09-18"
    t5 = next(t for t in ex["tml"] if t["point"] == "TML-05")
    assert t5["current_mm"] == 6.1 and 0 < t5["region"]["y"] < 1 and t5["region"]["w"] > 0
    draft = scan_findings.detect(ex["fields"], ex["tml"], "E-310")
    assert draft["severity"] == "high" and "6.1 mm below minimum 6.5 mm at TML-05" in draft["title"]


def test_finding_rules_for_valve_certificates():
    from app.scan_findings import detect
    reg = {"x": 0, "y": 0, "w": 1, "h": 1}
    f = lambda v: {"value": v, "confidence": 0.99, "region": reg, "page": 1}  # noqa: E731
    assert detect({"set_pressure_barg": f("18.5"), "pop_pressure_barg": f("19.5")}, [], "PSV-1")["severity"] == "high"
    assert detect({"next_due": f("2026-10-04")}, [], "PSV-1", today=date(2026, 9, 20))["severity"] == "medium"
    assert detect({"next_due": f("2026-10-04")}, [], "PSV-1", today=date(2026, 6, 1)) is None
    assert detect({"seat_leakage": f("FAILED - leaking")}, [], "PSV-1")["title"].endswith("seat leakage test failed")


def test_scan_to_finding_with_review_then_lane_routing_and_note(login):
    hse = login("ramesh.pm")
    t = hse.post(f"/api/documents/{_doc(login, 'DV-CAL-25-118')['id']}/scan-findings").json()
    r = t["result"]
    assert r["tag"] == "PSV-118" and any(c["result"] == "MISMATCH" and c["check"] == "Make / model" for c in r["checks"])
    if not r.get("finding"):
        pytest.skip("certificate not yet due on this date")
    assert t["status"] == "waiting_review" and r["finding"]["review_items"] >= 1
    fid = r["finding"]["id"]
    assert all(f["id"] != fid for f in login("contractor.x").get("/api/findings").json())
    items = [i for i in login("suresh.em").get("/api/review-items").json() if i["finding_id"] == fid]
    assert items and items[0]["region"]["w"] > 0
    assert login("ravi.e").post(f"/api/review-items/{items[0]['id']}", json={"resolution": "confirm"}).status_code == 404
    out = login("suresh.em").post(f"/api/review-items/{items[0]['id']}", json={"resolution": "confirm"}).json()
    assert out["state"] == "PENDING" and out["note"]["file_name"].endswith(".docx")
    f = next(x for x in login("suresh.em").get("/api/findings").json() if x["id"] == fid)
    assert f["lane"] == "electrical" and f["note_url"] and f["evidence_region"]["page"] == 1
    note = login("suresh.em").get(f["note_url"])
    assert note.status_code == 200 and note.content[:2] == b"PK"  # a real .docx
    assert login("ravi.e").get(f["note_url"]).status_code == 404
    done = login("ramesh.pm").get(f"/api/tasks/{t['id']}").json()
    assert done["status"] == "done" and any(s["name"] == "Draft approval note" and s["status"] == "done" for s in done["steps"])


def test_scanning_a_report_twice_does_not_duplicate(login):
    t = login("ramesh.pm").post(f"/api/documents/{_doc(login, 'UT-E310-2026-07')['id']}/scan-findings").json()
    assert t["result"]["finding"]["duplicate"] is True
    assert login("arjun.pl").post(f"/api/documents/{_doc(login, 'UT-E310-2026-07')['id']}/scan-findings").status_code in (403, 404)


# ------------------------------------------------------------------ agentic workbench and sandbox
def test_calculation_with_steps_delivers_an_excel_file(login):
    u = login("ravi.e")
    t = u.post("/api/work", json={"deliverable": "calc", "goal": "A2 motor full-load current",
                                  "calculation": "P = 75 * 1000 [W]\nI = P / (sqrt(3) * V * pf * eff) [A]",
                                  "variables": {"V": 415, "pf": 0.86, "eff": 0.94}}).json()
    assert t["status"] == "done", t
    steps = t["result"]["calculation"]["steps"]
    assert steps[1]["substituted"].startswith("75000 / (sqrt(3) * 415 * 0.86 * 0.94)") and abs(steps[1]["value"] - 129.0) < 1
    f = t["files"][0]
    r = u.get(f["url"])
    assert r.status_code == 200 and r.content[:2] == b"PK" and f["file_name"].endswith(".xlsx")
    assert login("meera.me").get(f["url"]).status_code == 404


def test_word_and_powerpoint_deliverables_are_cited_and_policy_filtered(login):
    u = login("ravi.e")
    for kind, ext in (("docx", ".docx"), ("pptx", ".pptx"), ("approval_note", ".docx")):
        t = u.post("/api/work", json={"deliverable": kind, "goal": "Summary of A2 trip troubleshooting and isolation"}).json()
        assert t["status"] == "done", t
        assert t["files"][0]["file_name"].endswith(ext) and u.get(t["files"][0]["url"]).content[:2] == b"PK"
        assert all(s["department"] != "Finance & Accounts" for s in t["result"]["sources"])
        names = [s["name"] for s in t["steps"]]
        assert names[:2] == ["Authorize request", "search_documents"] and "draft_text" in names
    import docx
    from app.db import q1
    path = q1("SELECT file_path FROM reports WHERE id=?", (t["files"][0]["id"],))["file_path"]
    text = "\n".join(p.text for p in docx.Document(path).paragraphs)
    assert "Approval route" in text and "Sign-off" in text


def test_excel_export_of_asset_records(login):
    t = login("ravi.e").post("/api/work", json={"deliverable": "xlsx", "goal": "export A2 work orders and ledger"}).json()
    assert t["status"] == "done", t
    import openpyxl
    from app.db import q1
    wb = openpyxl.load_workbook(q1("SELECT file_path FROM reports WHERE id=?", (t["files"][0]["id"],))["file_path"])
    assert "A2 work orders" in wb.sheetnames and "Sources" in wb.sheetnames


def test_sandbox_runs_code_without_network_or_system_access(login):
    u = login("ravi.e")
    t = u.post("/api/work", json={"deliverable": "code", "goal": "trend", "run": True,
                                  "code": "import statistics\nv=[7.9,7.1,6.1]\nprint(round(statistics.mean(v),2))\n"
                                          "open('trend.csv','w').write('t,v\\n1,7.9\\n')"}).json()
    assert t["status"] == "done" and t["result"]["run"]["stdout"].strip() == "7.03"
    assert t["files"][0]["file_name"] == "trend.csv"
    bad = u.post("/api/work", json={"deliverable": "code", "goal": "x", "run": True, "code": "import socket\nsocket.socket()"}).json()
    assert bad["result"]["run"]["exit_code"] != 0 and "not available" in bad["result"]["run"]["stderr"]
    bad = u.post("/api/work", json={"deliverable": "code", "goal": "x", "run": True, "code": "import os"}).json()
    assert "not available in the Yukti sandbox" in bad["result"]["run"]["stderr"]
    c = login("contractor.x").post("/api/work", json={"deliverable": "code", "goal": "x", "run": True, "code": "print(1)"}).json()
    assert c["status"] == "error" and "Contract workers" in c["error"]


def test_tools_cannot_write_to_control_systems():
    from app.policy import pdp, Subject
    s = Subject(id="x", username="x", display_name="x", post="", department="Operations (Production)", clearance=4,
                roles=["plant_manager"], asset_scopes=["*"], grants=[])
    assert not pdp.decide(s, "use", {"type": "tool", "tool": "dcs_write", "tool_class": "control_write"}).allowed


def test_calculate_rejects_code():
    from app.workbench import calculate
    with pytest.raises(ValueError):
        calculate("x = __import__('os').system('echo hi')")
    assert calculate("a = 2\nb = a ** 3")["result"]["value"] == 8


# ------------------------------------------------------------------ isolated research layer
def test_research_layer_is_isolated_and_permissioned(login):
    a = login("admin")
    assert a.put("/api/admin/research", json={"url": "http://10.1.2.3:8090/v1"}).status_code == 422  # must stay on this server
    assert a.put("/api/admin/research", json={"enabled": True, "guardrails": "Refinery research rules."}).json()["enabled"]
    assert login("ravi.e").post("/api/research/ask", json={"question": "hello"}).status_code == 403
    from app.policy import pdp
    from app.auth import subject_for
    from app.db import q1
    s = subject_for(q1("SELECT * FROM users WHERE username='ramesh.pm'"))
    doc = {"type": "document", "id": "d", "department": "Operations (Production)", "classification": 0, "doc_type": "SOP",
           "asset_tags": [], "units": [], "status": "CURRENT"}
    assert pdp.decide(s, "read", doc).allowed and not pdp.decide(s, "read", doc, {"layer": "research"}).allowed
    assert not pdp.decide(s, "use", {"type": "tool", "tool": "calculate", "tool_class": "compute"}, {"layer": "research"}).allowed


# ------------------------------------------------------------------ forecast + optimizer plan
def test_forecast_estimates_demand_with_an_interval():
    from app.forecast import forecast_series
    f = forecast_series([100, 104, 108, 113, 116, 121, 124, 129])
    assert f["method"] == "Holt linear smoothing" and f["low"] < f["expected"] < f["high"] and 128 < f["expected"] < 138
    assert forecast_series([10, 11]) is None and forecast_series([10, 12, 11])["method"].startswith("mean")


def test_production_plan_scenario_card(login):
    p = login("arjun.pl")
    m = p.get("/api/production/model").json()
    hist = [{"month": f"2026-{i:02d}", "kt": v} for i, v in enumerate([300, 310, 318, 330, 335, 348, 352, 360], 1)]
    body = {"crude_cost_usd_t": 520, "crude_max_kt": 1100, "notes": "test",
            "units": [{"code": "CDU", "name": "Crude unit", "capacity_kt_month": 1250, "feed": "crude", "enabled": True,
                       "planned_down_days": 0, "yields": {"Diesel": 0.42, "Petrol": 0.2, "Petchem": 0.18}}],
            "products": [{"name": "Diesel", "price_usd_t": 690, "demand_history": hist, "inventory_kt": 20, "storage_max_kt": 120},
                         {"name": "Petrol", "crack_usd_t": 140, "dispatch_max_kt": 200},
                         {"name": "Petchem", "price_usd_t": 610, "contract_min_kt": 50}]}
    assert p.put("/api/production/model", json=body).status_code == 200
    card = p.post("/api/production/plan", json={"petchem_margin_delta": -60}).json()
    assert "Diesel" in card["forecast"] and card["forecast"]["Diesel"]["low"] < card["forecast"]["Diesel"]["high"]
    assert {a["case"] for a in card["alternatives"]} == {"low", "high"}
    assert "margin_usd" in card["expected_value"] and card["assumptions"] and "planners decide" in card["decision"].lower()
    assert card["scenario"]["scenario"]["Diesel"] <= card["forecast"]["Diesel"]["expected"] + 20 + 1e-6  # demand cap + stock
    assert p.put("/api/production/model", json=m).status_code == 200


# ------------------------------------------------------------------ Performance Engine profiles
def test_profiles_are_selected_from_model_size_and_free_vram():
    from app.perf_profiles import PROFILES, apply, params_b, select
    assert params_b({"name": "gemma-3-27b-it-Q4_K_M"}) == 27 and params_b({"name": "gemma-3-4b-it"}) == 4
    assert select({"name": "gemma-3-4b-it"}, 6000) == "small-6gb"
    assert select({"name": "gemma-3-4b-it"}, None) == "auto-fit"  # GPU memory not readable: never fall back to CPU only
    assert select({"name": "gemma-3-4b-it"}, 2000) == "small-cpu"
    assert select({"name": "gemma-3-27b-it"}, 12000) == "reasoning-27b-12gb"
    assert select({"name": "gemma-3-27b-it"}, 24576) == "reasoning-27b-24gb"
    cfg = apply("small-6gb", {"ctx_size": 8192})
    assert cfg["ctx_size"] == 8192 and cfg["cache_type_k"] == "q8_0" and cfg["profile"] == "small-6gb"
    assert PROFILES["small-6gb"]["status"] == "measured" and PROFILES["reasoning-27b-16gb"]["status"] == "reference"


def test_profiles_api_is_admin_only(login):
    out = login("admin").get("/api/admin/profiles").json()
    assert {p["id"] for p in out["profiles"]} >= {"small-6gb", "reasoning-27b-12gb"}
    assert login("ravi.e").get("/api/admin/profiles").status_code == 403
    assert login("admin").post("/api/admin/profiles/small-6gb/validate").status_code == 409  # no model loaded in CI


def test_worker_models_and_kv_settings_are_validated(login):
    a = login("admin")
    s = a.get("/api/admin/ai-settings").json()
    s["worker_models"] = {"reasoning": {"model_id": "gemma-3-27b-it-Q4_K_M"}}
    s["kv_ttl_s"] = 900
    out = a.put("/api/admin/ai-settings", json=s).json()
    assert out["worker_models"]["reasoning"]["model_id"] and out["kv_ttl_s"] == 900
    s["worker_models"] = {"nonsense": {"model_id": "x"}}
    assert a.put("/api/admin/ai-settings", json=s).status_code == 422


# ------------------------------------------------------------------ speech: AI4Bharat models are optional add-ons
def test_ai4bharat_models_are_used_only_when_installed():
    from app import speech
    assert set(speech.INDIC) == {"hi", "kn"} and speech.indic_model("en") is None
    assert all(n.startswith("AI4Bharat") for _, n in speech.INDIC.values())


# ------------------------------------------------------------------ demonstration mode on a production-started database
def test_demo_mode_relists_only_untouched_sample_accounts(app_client):
    from app import seed
    from app.auth import hash_password
    from app.db import ex, q1
    ex("UPDATE users SET demo_password=NULL, must_change_password=1 WHERE username IN ('anil.u', 'meera.me')")
    ex("UPDATE users SET password_hash=? WHERE username='meera.me'", (hash_password("Changed!Pass-2026"),))
    seed.ensure_demo_accounts()
    assert q1("SELECT demo_password FROM users WHERE username='anil.u'")["demo_password"] == "Anil@2026"
    assert q1("SELECT demo_password FROM users WHERE username='meera.me'")["demo_password"] is None  # changed: stays private
    ex("UPDATE users SET password_hash=?, demo_password='Meera@2026', must_change_password=0 WHERE username='meera.me'",
       (hash_password("Meera@2026"),))
